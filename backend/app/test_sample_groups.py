"""The Step 2 build list's defining-SNP groups are the groups vsnp3 would make.

The list shows which groups each VCF falls into, and its filter selects by
them, so a comparison can be chosen by group before it is run. A group shown
there that the run then disagrees with would choose the wrong samples, so the
reader (app/sample_groups.py) is held to vsnp3's own code:

  * where vsnp3 is importable (the app's env), its VCF normalizer and reader
    (VCF_to_DF._write_normalized / read_vcf) and its Group.group_selection are
    run on the same VCFs, and every sample's groups must agree, at default
    thresholds, a raised QUAL, a raised MQ, with filters off and with -hash;
  * the VCFs are built to hit each rule: QUAL and MQ truncation at the
    threshold, AC=1 and AC="1,1", an INFO item with two "=", ";MQM=",
    quoted and space-separated lines, a duplicate position whose first record
    is an indel, a zero-coverage record at a defining position, a leading-zero
    POS, a gzipped VCF, CRLF line ends, a sample with no good SNP at all
    against a "!"-only group, a two-position group, a filtered position and
    a "#" column;
  * the cache: a second look reads nothing, a touched file is read again, a
    changed threshold reads nothing, changed defining positions read all;
  * a later look costs one stat per VCF;
  * /step2/sample-groups and /name-aliases on a project.

Validated beyond this by hand on 2026-09-29 against two real vsnp3 runs of the
mtbc0_v1.1 test set: 68 of 68 samples at default settings, 20 of 20 at
-w 1000 -n, every group exact.

Run directly:  <env>/bin/python backend/app/test_sample_groups.py
"""
from __future__ import annotations

import gzip
import os
import sys
import tempfile
import threading
import types
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import openpyxl  # noqa: E402

import app.main as m  # noqa: E402
from app import sample_groups as sg  # noqa: E402

FAILED = 0


def check(label, cond):
    global FAILED
    if cond:
        print(f"  OK  {label}")
    else:
        FAILED += 1
        print(f"  FAIL {label}")


class Calls:
    """Counts os.stat / os.lstat / os.scandir / os.listdir calls."""

    def __init__(self):
        self.n = 0
        self._lock = threading.Lock()
        self._saved = {}

    def __enter__(self):
        for name in ("stat", "lstat", "scandir", "listdir"):
            real = getattr(os, name)
            self._saved[name] = real

            def counted(*a, _real=real, **k):
                with self._lock:
                    self.n += 1
                return _real(*a, **k)
            setattr(os, name, counted)
        return self

    def __exit__(self, *exc):
        for name, real in self._saved.items():
            setattr(os, name, real)


# --- Fixtures ----------------------------------------------------------------

HEADER = ("##fileformat=VCFv4.2\n##contig=<ID=chr1,length=500>\n"
          "#CHROM\tPOS\tID\tREF\tALT\tQUAL\tFILTER\tINFO\tFORMAT\tS\n")


def rec(pos, qual="500", info="AC=2;DP=40;MQ=60", ref="A", alt="G"):
    return f"chr1\t{pos}\t.\t{ref}\t{alt}\t{qual}\t.\t{info}\tGT\t1/1\n"


ZC = "chr1\t{}\t.\tN\t.\t.\t.\t.\tGT\t./.\n"

# Positions: 10 Group-A, 20 Group-A1 (a child, 10 AND 20), 30 Group-B,
# 50 "!" alone (Group-notC), 60 filtered by the first column, 70 "#"-marked.
VCFS = {
    "good_a": rec(10) + rec(20),
    "good_b": rec(30),
    "qual_at_threshold": rec(10, qual="150.9") + rec(30, qual="151"),
    "mq_below": rec(10, info="AC=2;MQ=55.9") + rec(30, info="AC=2;MQ=56"),
    "het_and_multi_ac": rec(10, info="AC=1;MQ=60") + rec(30, info="AC=1,1;MQ=60") + rec(20, info="AC=2.0;MQ=60"),
    "info_two_equals": rec(10, info="ANN=x=y;AC=2;MQ=60") + rec(30),
    "mqm": rec(10, info="AC=2;MQM=60") + rec(30, info="MQM=60;AC=2") + rec(99),
    "quoted_and_spaced": '"' + rec(10).rstrip("\n") + '"\n' + rec(30).replace("\t", "  ", 3),
    "dup_indel_first": rec(10, ref="AT", alt="A") + rec(10) + rec(30) + rec(30, qual="1"),
    "zc_at_defining": ZC.format(10) + ZC.format(30) + rec(99),
    "leading_zero_pos": rec("010", ref="AT", alt="A") + rec(10) + rec(30),
    "no_good_snp": rec(10, qual="20") + rec(99, info="AC=1;MQ=60"),
    "only_elsewhere": rec(99),
    "inverted_hit": rec(50) + rec(99),
    "filtered_pos": rec(60) + rec(10),
    "hashed_pos": rec(70) + rec(10),
    "crlf": (rec(10) + rec(20)).replace("\n", "\r\n"),
}


def write_define_filter(path: Path) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    cols = [(None, "All"), ("chr1:10", "Group-A"), ("chr1:10, chr1:20", "Group-A1"), ("chr1:30", "Group-B"),
            ("chr1:50!", "Group-notC"), ("chr1:60", "Group-D"), ("#chr1:70", "Group-E")]
    for c, (hdr, grp) in enumerate(cols, start=1):
        ws.cell(row=1, column=c, value=hdr)
        ws.cell(row=2, column=c, value=grp)
    ws.cell(row=3, column=1, value="chr1:59-61")     # a filter range in the "all" column
    wb.save(path)
    return path


def write_vcfs(db: Path) -> list:
    db.mkdir(parents=True, exist_ok=True)
    files = []
    for name, body in VCFS.items():
        p = db / f"{name}_zc.vcf"
        p.write_bytes((HEADER + body).encode())
        files.append((p.name, p))
    gz = db / "gzipped_zc.vcf.gz"
    with gzip.open(gz, "wb") as fh:
        fh.write((HEADER + rec(30)).encode())
    files.append((gz.name, gz))
    return files


# --- vsnp3's own code, as the oracle -----------------------------------------

def vsnp3_oracle():
    """(read(path) -> df or None, groups(dfs, specs) -> {sample: [groups]}) from
    the installed vsnp3, or None when it is not importable here."""
    env_bin = Path(__file__).resolve().parents[2] / "env" / "bin"
    if not (env_bin / "vsnp3_step2.py").exists():
        return None
    sys.path.insert(0, str(env_bin))
    try:
        import vsnp3_step2 as s2
        import vsnp3_group_on_defining_snps as gods
    except Exception as exc:                   # not this env's vsnp3
        print(f"  (vsnp3 not importable: {exc})")
        return None
    import pandas as pd
    norm_dir = Path(tempfile.mkdtemp(prefix="vsnp3_norm_"))
    stub = types.SimpleNamespace(first_sample_only=True, assume_gt_only_quality=False)

    def read(path: Path):
        src = path
        if str(path).endswith(".gz"):
            src = norm_dir / (path.name[:-3] + ".raw")
            src.write_bytes(gzip.decompress(path.read_bytes()))
        out = norm_dir / path.name.replace(".gz", "")
        s2.VCF_to_DF._write_normalized(stub, str(src), str(out))
        try:
            df = s2.VCF_to_DF.read_vcf(None, str(out))
        except ValueError:
            return None                         # no MQ anywhere: vsnp3 drops the sample
        df["abs_pos"] = df["CHROM"] + ":" + df["POS"].astype(str)
        return df.drop_duplicates(subset=["abs_pos"])

    def groups(dfs, defs, qual, mq, no_filters):
        # The filter list as Group.__init__ builds it, expanded by vsnp3 itself.
        filter_all = gods.Group.list_expansion(None, defs.filter_raw)
        essentials = {}
        for name, df in dfs.items():
            if df is None:
                continue
            if not no_filters and filter_all:
                df = df[~df["abs_pos"].isin(filter_all)]
            df = df[df["QUAL"].notna() & (df["QUAL"] > qual) & (df["AC"] == 2)
                    & (df["REF"].str.len() == 1) & (df["ALT"].str.len() == 1)
                    & df["MQ"].notna() & (df["MQ"] >= mq)]
            if not df.empty:
                essentials[name] = df
        grp = object.__new__(gods.Group)
        grp.dataframe_essentials = essentials
        out = {name: [] for name in dfs}
        spec = defs.spec_strings
        for abs_pos, gname in spec:
            found, sample_dict = gods.Group.group_selection(grp, abs_pos)
            if found:
                for s in sample_dict:
                    out[s].append(gname)
        return out

    return read, groups, pd


def spec_strings(path: Path, hash_groups: bool):
    """The (column header, group) pairs vsnp3's Group.__init__ builds, and the
    first column's raw filter entries."""
    import pandas as pd
    if hash_groups:
        df = pd.read_excel(path, header=None)
        first = df.iloc[0]
        df.iloc[0] = first.apply(lambda x: str(x).replace("#", "") if pd.notna(x) else x)
        df.columns = df.iloc[0]
        df = df.drop(df.index[0]).reset_index(drop=True)
        df = df.dropna(axis=1, how="all")
        df = df.loc[:, (df != "").any()]
    else:
        df = pd.read_excel(path)
    pairs = list(df.iloc[:, 1:].head(n=1).to_dict(orient="records")[0].items())
    raw = [x for x in df.iloc[:, 0].to_list()[1:] if str(x) != "nan"]
    return pairs, raw


def main() -> int:
    tmp = Path(tempfile.mkdtemp(prefix="sample_groups_"))
    ref_dir = tmp / "refs" / "H5"
    ref_dir.mkdir(parents=True)
    dfx = write_define_filter(ref_dir / "H5_define_filter.xlsx")
    db = tmp / "projects" / "p1" / "step2" / "vcf_database"
    files = write_vcfs(db)

    print("[the defining-SNP workbook, read as vsnp3 reads it]")
    defs = sg.load_definitions(dfx)
    check("every column is a group, in workbook order",
          [g[0] for g in defs.groups] == ["Group-A", "Group-A1", "Group-B", "Group-notC", "Group-D", "Group-E"])
    check("...the '#' column's position is spelled with its '#', so it never matches",
          ("Group-E", ("#chr1:70",), ()) in defs.groups)
    check("a two-position column needs both, a '!' position is inverted",
          ("Group-A1", ("chr1:10", "chr1:20"), ()) in defs.groups
          and ("Group-notC", (), ("chr1:50",)) in defs.groups)
    check("the first column's range is expanded into the filter set",
          {"chr1:59", "chr1:60", "chr1:61"} <= defs.filter_all)
    hdefs = sg.load_definitions(dfx, hash_groups=True)
    check("with -hash the '#' is dropped and the column is live",
          ("Group-E", ("chr1:70",), ()) in hdefs.groups and defs.digest != hdefs.digest)

    print("[the reader against hand-written expectations, default settings]")
    out = sg.sample_groups(files, defs, None)
    g = out["groups"]
    want = {
        "good_a_zc.vcf": ["Group-A", "Group-A1", "Group-notC"],
        "good_b_zc.vcf": ["Group-B", "Group-notC"],
        "qual_at_threshold_zc.vcf": ["Group-B", "Group-notC"],        # 150.9 truncates to 150
        "mq_below_zc.vcf": ["Group-B", "Group-notC"],                 # 55.9 truncates to 55
        "het_and_multi_ac_zc.vcf": ["Group-notC"],                    # AC=2.0 is 2; 10 is AC=1
        "info_two_equals_zc.vcf": ["Group-B", "Group-notC"],          # the bad item voids AC and MQ
        "mqm_zc.vcf": ["Group-A", "Group-notC"],                      # ";MQM=" renamed, leading MQM not
        "quoted_and_spaced_zc.vcf": ["Group-A", "Group-B", "Group-notC"],
        "dup_indel_first_zc.vcf": ["Group-B", "Group-notC"],          # first record at 10 is the indel
        "zc_at_defining_zc.vcf": ["Group-notC"],
        "leading_zero_pos_zc.vcf": ["Group-B", "Group-notC"],         # "010" is 10, and it came first
        "no_good_snp_zc.vcf": [],                                     # not even the "!" group
        "only_elsewhere_zc.vcf": ["Group-notC"],
        "inverted_hit_zc.vcf": [],
        "filtered_pos_zc.vcf": ["Group-A", "Group-notC"],             # 60 is filtered out
        "hashed_pos_zc.vcf": ["Group-A", "Group-notC"],
        "crlf_zc.vcf": ["Group-A", "Group-A1", "Group-notC"],
        "gzipped_zc.vcf.gz": ["Group-B", "Group-notC"],
    }
    for name, groups in want.items():
        check(f"{name}: {groups}", g.get(name) == groups)
    check("with filters off, the filtered position counts",
          sg.sample_groups(files, defs, None, no_filters=True)["groups"]["filtered_pos_zc.vcf"]
          == ["Group-A", "Group-notC", "Group-D"])

    oracle = vsnp3_oracle()
    if oracle is None:
        print("[vsnp3 oracle skipped: vsnp3 is not importable from this checkout's env]")
    else:
        read, oracle_groups, _pd = oracle
        dfs = {name: read(path) for name, path in files}
        settings = [
            ("defaults", {}, False),
            ("-w 400", {"qual_threshold": 400}, False),
            ("-y 61", {"mq_threshold": 61}, False),
            ("-n", {"no_filters": True}, False),
            ("-hash", {}, True),
        ]
        print("[the reader against vsnp3's own code, setting by setting]")
        for label, kw, hashed in settings:
            d = sg.load_definitions(dfx, hash_groups=hashed)
            pairs, raw = spec_strings(dfx, hashed)
            d_or = types.SimpleNamespace(spec_strings=pairs, filter_raw=raw)
            mine = sg.sample_groups(files, d, None, **kw)["groups"]
            theirs = oracle_groups(dfs, d_or, kw.get("qual_threshold", 150), kw.get("mq_threshold", 56),
                                   kw.get("no_filters", False))
            differ = sorted(n for n in theirs if sorted(theirs[n]) != sorted(mine.get(n, [])))
            check(f"{label}: every one of {len(theirs)} samples in the groups vsnp3 gives"
                  + (f" (differ: {differ})" if differ else ""), not differ)

    print("[the cache]")
    cache = db.parent / sg.CACHE_BASENAME
    first = sg.sample_groups(files, defs, cache)
    again = sg.sample_groups(files, defs, cache)
    check("the first look reads every VCF, the second none, with the same answer",
          first["read"] == len(files) and again["read"] == 0 and first["groups"] == again["groups"])
    raised = sg.sample_groups(files, defs, cache, qual_threshold=1000)
    check("a changed threshold is answered from the cache, and changes the answer",
          raised["read"] == 0 and raised["groups"]["good_a_zc.vcf"] == [])
    p = dict(files)["good_b_zc.vcf"]
    p.write_bytes((HEADER + rec(10)).encode())
    os.utime(p, ns=(p.stat().st_atime_ns, p.stat().st_mtime_ns + 5_000_000_000))
    touched = sg.sample_groups(files, defs, cache)
    check("an edited VCF is read again, and only it",
          touched["read"] == 1 and touched["groups"]["good_b_zc.vcf"] == ["Group-A", "Group-notC"])
    moved = sg.load_definitions(write_define_filter(ref_dir / "H5_define_filter.xlsx"))
    wb = openpyxl.load_workbook(dfx)
    wb.active.cell(row=1, column=4, value="chr1:31")
    wb.save(dfx)
    os.utime(dfx, ns=(dfx.stat().st_atime_ns, dfx.stat().st_mtime_ns + 5_000_000_000))
    redefined = sg.load_definitions(dfx)
    check("moving a defining position changes the definitions' stamp",
          redefined.digest != moved.digest)
    check("...and every VCF is read again",
          sg.sample_groups(files, redefined, cache)["read"] == len(files))
    write_define_filter(dfx)
    os.utime(dfx, ns=(dfx.stat().st_atime_ns, dfx.stat().st_mtime_ns + 10_000_000_000))
    defs = sg.load_definitions(dfx)

    print("[bounded requests]")
    many = [(f"copy{i}_zc.vcf", dict(files)["good_a_zc.vcf"]) for i in range(600)]
    part = sg.sample_groups(many, defs, None, budget_s=0)
    check("a request always reads one batch, even with no time left, and says what is pending",
          part["read"] == 256 and part["pending"] == 600 - 256)

    print("[a later look is one stat per VCF]")
    sg.sample_groups(files, defs, cache)
    with Calls() as c:
        warm = sg.sample_groups(files, defs, cache)
    check(f"{len(files)} VCFs, {c.n} calls, nothing read", warm["read"] == 0 and c.n == len(files))

    print("[which workbook is the defining-SNP one: vsnp3's rule]")
    check("the define_filter, beside remove, metadata and a lock file",
          (ref_dir / "H5_remove_from_analysis.xlsx").write_bytes(b"") is not None
          and (ref_dir / "H5_metadata.xlsx").write_bytes(b"") is not None
          and (ref_dir / "~$H5_define_filter.xlsx").write_bytes(b"") is not None
          and sg.define_filter_of(ref_dir) == (dfx, ""))
    (ref_dir / "copy of defining.xlsx").write_bytes(b"")
    path, why = sg.define_filter_of(ref_dir)
    check("two candidates is vsnp3's refusal, named", path is None and "more than one" in why)
    (ref_dir / "copy of defining.xlsx").unlink()
    check("no reference folder says so", sg.define_filter_of(None)[0] is None)

    print("[the endpoints]")
    proj = tmp / "projects" / "p1"
    (proj / "project.json").write_text('{"reference": "H5"}', encoding="utf-8")
    m.load_config = lambda: {"vsnp3_path": str(tmp / "vsnp3"), "projects_root": str(proj.parent)}
    m._project_dir_for = lambda _cfg, _name: proj
    m.reference_roots = lambda _p: [ref_dir.parent]
    got = m.step2_sample_groups("p1")
    check("/step2/sample-groups answers for every VCF in vcf_database",
          got["available"] and got["total"] == len(files) and got["pending"] == 0
          and got["groups"]["good_a_zc.vcf"] == ["Group-A", "Group-A1", "Group-notC"])
    check("...at the thresholds it is given",
          m.step2_sample_groups("p1", qual_threshold=1000)["groups"]["good_a_zc.vcf"] == [])
    m.reference_roots = lambda _p: []
    gone = m.step2_sample_groups("p1")
    check("no reference folder: unavailable, with the reason", not gone["available"] and gone["reason"])
    m.reference_roots = lambda _p: [ref_dir.parent]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.cell(row=1, column=1, value="good_a_zc.vcf")         # column A names the FILE
    ws.cell(row=1, column=2, value="Chicken NE/2022")
    wb.save(ref_dir / "H5_metadata.xlsx")
    na = m.project_name_aliases("p1")
    check("/name-aliases gives the metadata name vsnp3 prints, keyed by the sample",
          na["display"].get("good_a") == "Chicken_NE_2022")
    check("...only for samples the metadata names", "good_b" not in na["display"])

    print("PASS" if not FAILED else f"{FAILED} FAILURE(S)")
    return 1 if FAILED else 0


if __name__ == "__main__":
    raise SystemExit(main())
