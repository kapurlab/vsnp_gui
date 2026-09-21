"""SNP-table rows link to Step 1 when the reference metadata RENAMED the sample.

Reported case: Step 1 ran `EPI-ISL-19152935`, and its BAM sits on disk under
that name. The Step 2 SNP table calls the same specimen
`GISAID-19152935_FAV-0863-5_Razorbill_2022_CAN-NL`, because the reference's
`*_metadata.xlsx` says to. Clicking one of that row's SNPs did nothing: no IGV,
no hover, no explanation.

`_canonical_stem`'s four rules all compare the two strings, and none of them can
match these two — the metadata replaced the id (`EPI-ISL-` to `GISAID-`) rather
than decorating it. The mapping is not recoverable from the names; it is written
down in the workbook vSNP3 renamed from, and `name_aliases` reads it.

What is pinned here: the alias resolves the row in both renderers, it does not
disturb a table whose names never needed it, an ambiguous alias is refused
rather than guessed, and — the question that decides whether this can ship to a
40,000-sample database — consulting the index makes the resolution cheaper, not
dearer, because a hit returns before the per-row scan over every known sample.

Run directly:  <conda>/bin/python backend/app/test_metadata_aliases.py
"""
from __future__ import annotations

import re
import sys
import tempfile
import time
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

import openpyxl  # noqa: E402

import name_aliases  # noqa: E402
import xlsx_html  # noqa: E402
from test_xlsx_filter import make_group_table  # noqa: E402

FAILED = 0


def check(label, cond):
    global FAILED
    if cond:
        print(f"  OK  {label}")
    else:
        FAILED += 1
        print(f"  FAIL {label}")


def write_metadata(path: Path, pairs) -> Path:
    """A reference metadata workbook: column A stored name, B display name."""
    wb = openpyxl.Workbook()
    ws = wb.active
    for i, (orig, disp) in enumerate(pairs, start=1):
        ws.cell(row=i, column=1, value=orig)
        ws.cell(row=i, column=2, value=disp)
    wb.save(path)
    wb.close()
    return path


def linked_rows(html_rows) -> int:
    """Rows carrying at least one clickable variant cell."""
    return sum(1 for r in html_rows if "xlsx-variant" in r)


def render(book, n_labels, cols, bams, aliases):
    return xlsx_html.render_window(
        book, n_labels + 4, cols, None, "proj",
        samples_with_bams=bams, samples_with_vcfs=set(),
        max_cells=1_000_000, max_rows=1_000,
        max_table_bytes=64 * 1024 * 1024, aliases=aliases,
    )


def main() -> int:
    tmp = Path(tempfile.mkdtemp(prefix="meta_alias_"))

    print("[the reported case: metadata replaced the id]")
    on_disk = "EPI-ISL-19152935"
    in_table = "GISAID-19152935_FAV-0863-5_Razorbill_2022_CAN-NL"
    meta = write_metadata(tmp / "H5_metadata.xlsx", [(on_disk, in_table)])
    idx = name_aliases.load(meta)
    check("the workbook was read", len(idx) == 1)
    check("display -> stored", idx.original_of(in_table) == on_disk)
    check("stored -> display", idx.display_of(on_disk) == in_table)
    check("no string rule could have done it",
          xlsx_html._canonical_stem(in_table, {on_disk}) == in_table)
    check("with the index, it resolves",
          xlsx_html._canonical_stem(in_table, {on_disk}, aliases=idx) == on_disk)

    print("[found by the directory, the way a request finds it]")
    check("metadata_file picks the workbook out of a reference dir",
          name_aliases.metadata_file(tmp) == meta)
    check("a reference with no metadata is empty, not an error",
          not name_aliases.for_reference_dir(tmp / "nope"))

    print("[rendered: the row's SNP cells become IGV links]")
    book = tmp / "group.xlsx"
    make_group_table(book, [in_table], 10, lambda i, c: "A" if c >= 3 else "C")
    bams = {on_disk}
    without = render(book, 1, 10, bams, None)
    with_idx = render(book, 1, 10, bams, idx)
    check("unlinked without the index", linked_rows(without["rows"]) == 0)
    check("linked with it", linked_rows(with_idx["rows"]) == 1)
    check("...and the row's IGV target is the STEP 1 sample, not the label",
          with_idx["row_samples"] == ["", "", on_disk, "", ""])

    print("[the clade-filtered renderer agrees with the plain one]")
    filt = xlsx_html.render_filtered_window(
        book, 5, 10, None, "proj", bams, set(), [in_table],
        1_000_000, 1_000, 64 * 1024 * 1024, aliases=idx)
    check("the tip name selects its row", linked_rows(filt["rows"]) == 1)

    print("[a table that never needed the index is untouched]")
    plain_labels = ["SRR33643035", "19-1234"]
    plain_book = tmp / "plain.xlsx"
    make_group_table(plain_book, plain_labels, 10, lambda i, c: "A" if c >= 3 else "C")
    plain_bams = set(plain_labels)
    a = render(plain_book, 2, 10, plain_bams, None)
    b = render(plain_book, 2, 10, plain_bams, idx)
    check("identical output with and without", a["rows"] == b["rows"])

    print("[an ambiguous alias is refused, not guessed]")
    amb = write_metadata(tmp / "amb_metadata.xlsx", [
        ("sample-one", "Duck_2024_AK"),
        ("sample-two", "duck-2024-ak"),   # same string once case/sep folded
    ])
    aidx = name_aliases.load(amb)
    check("exact spellings still resolve",
          aidx.original_of("Duck_2024_AK") == "sample-one")
    check("a folded collision resolves to nothing",
          aidx.original_of("DUCK-2024-AK") is None)

    print("[the stored name may itself need the old rules]")
    # Step 1 stages with a suffix of its own; the metadata knows the bare id.
    staged = write_metadata(tmp / "staged_metadata.xlsx",
                            [("24-029315-007", "GISAID-777_GWTE_2024-09-26")])
    sidx = name_aliases.load(staged)
    check("resolves through to the staged folder name",
          xlsx_html._canonical_stem("GISAID-777_GWTE_2024-09-26",
                                    {"24-029315-007-original"},
                                    aliases=sidx) == "24-029315-007-original")

    print("[cost at database scale: 40,000 samples]")
    n = 40_000
    pairs = [(f"EPI-ISL-{19000000 + i}",
              f"GISAID-{19000000 + i}_SPECIES_{i % 50}_2024_REGION") for i in range(n)]
    big_meta = write_metadata(tmp / "big_metadata.xlsx", pairs)
    t0 = time.monotonic()
    big = name_aliases.load(big_meta)
    t_read = time.monotonic() - t0
    check(f"read {len(big):,} rows once in {t_read:.2f}s", len(big) == n)
    t0 = time.monotonic()
    name_aliases.load(big_meta)
    t_cached = time.monotonic() - t0
    check(f"a second request pays {t_cached * 1000:.2f}ms (cached)", t_cached < 0.01)

    known = {orig for orig, _ in pairs}
    labels = [disp for _, disp in pairs[:2000]]
    t0 = time.monotonic()
    for lab in labels:
        xlsx_html._canonical_stem(lab, known, aliases=big)
    t_with = time.monotonic() - t0
    # The same 2,000 rows with no index: every one falls through to the scans.
    t0 = time.monotonic()
    for lab in labels[:20]:
        xlsx_html._canonical_stem(lab, known)
    t_without_20 = time.monotonic() - t0
    t_without = t_without_20 * (len(labels) / 20)
    print(f"      2,000 rows against 40,000 samples: "
          f"{t_with * 1000:.0f} ms with the index, "
          f"~{t_without * 1000:.0f} ms without (extrapolated from 20 rows)")
    check("resolving with the index is faster, not slower", t_with < t_without)
    check("...and every row resolved",
          all(xlsx_html._canonical_stem(l, known, aliases=big) in known
              for l in labels[:200]))

    print("[vsnp3's own name cleaning is mirrored, so the PRINTED name resolves]")
    # vsnp3 rewrites column B before it reaches any table: accents folded,
    # space : , ( ) ; [ ] ' / . * ? { } become "_", runs of "_" collapse, and a
    # trailing "_" or "-" is dropped. What is on screen is the cleaned form.
    raw_b = "GISAID-19152935 FAV-0863-5 Razorbill 2022 CAN/NL."
    shown = "GISAID-19152935_FAV-0863-5_Razorbill_2022_CAN_NL"
    check("cleaning matches vsnp3", name_aliases.vsnp3_display_name(raw_b) == shown)
    check("accents fold to ASCII, like vsnp3",
          name_aliases.vsnp3_display_name("Qu\u00e9bec (2024)") == "Quebec_2024")
    cleaned = write_metadata(tmp / "cleaned_metadata.xlsx",
                             [("EPI-ISL-19152935_zc.vcf", raw_b)])
    cidx = name_aliases.load(cleaned)
    check("the name as PRINTED resolves to the stored sample",
          cidx.original_of(shown) == "EPI-ISL-19152935")
    check("the name as WRITTEN in the workbook resolves too",
          cidx.original_of(raw_b) == "EPI-ISL-19152935")
    check("display_of answers with the printed form, not the cell",
          cidx.display_of("EPI-ISL-19152935") == shown)
    check("...and a SNP-table label carrying _zc.vcf still resolves",
          cidx.original_of(shown + "_zc.vcf") == "EPI-ISL-19152935")

    print("[column A is matched the way vsnp3 matches it]")
    # vsnp3 compares the VCF basename to column A exactly, then with .vcf
    # stripped, then _zc, then _zc_*; every one of those spellings must reach
    # the same sample.
    check("file keys, cumulative", name_aliases.vsnp3_file_keys("S1_zc_L001.vcf.gz")
          == ["S1_zc_L001.vcf", "S1_zc_L001", "S1"])
    colA = write_metadata(tmp / "colA_metadata.xlsx",
                          [("EPI-ISL-19152935", "GISAID-19152935_X"),   # bare id in column A
                           ("SRR1791698_zc.vcf", "SRR1791698_updated")])
    aidx2 = name_aliases.load(colA)
    check("a bare column A still maps the on-disk stem both ways",
          aidx2.display_of("EPI-ISL-19152935_zc.vcf") == "GISAID-19152935_X"
          and aidx2.original_of("GISAID-19152935_X") == "EPI-ISL-19152935")
    check("a suffixed column A answers with the stem",
          aidx2.original_of("SRR1791698_updated") == "SRR1791698")

    print("[the endpoints the two panes call]")
    sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
    from app import main as M

    ref_dir = tmp / "refs/H5N1"
    ref_dir.mkdir(parents=True)
    (ref_dir / "H5_define_filter.xlsx").write_bytes(b"")
    write_metadata(ref_dir / "H5_metadata.xlsx", [
        # The shape a real workbook has: column A names the VCF FILE.
        (f"{on_disk}_zc.vcf", in_table),
        ("SRR1791698_zc.vcf", "SRR1791698_updated_name"),
    ])
    proj = tmp / "projects/p1"
    proj.mkdir(parents=True)
    (proj / "project.json").write_text('{"reference": "H5N1"}', encoding="utf-8")

    M.load_config = lambda: {"vsnp3_path": "/nonexistent", "projects_root": str(proj.parent)}
    M.list_references = lambda _p: [
        {"name": "H5N1", "path": str(ref_dir), "root": str(ref_dir.parent), "shadowed": []}]
    # A project's reference dir is resolved by NAME against the registered
    # roots (vsnp3's own `-t` rule), not by surveying every reference — see
    # _project_reference_dir.
    M.reference_roots = lambda _p: [ref_dir.parent]
    M._project_dir_for = lambda _cfg, _name: proj

    res = M.resolve_names("p1", M.NameResolveRequest(names=[in_table, on_disk, "not-a-sample"]))
    check("the SNP-table spelling resolves to the stored name",
          res["counterparts"].get(in_table) == [on_disk])
    check("...and the stored name back to the table spelling",
          res["counterparts"].get(on_disk) == [in_table])
    check("a name the metadata does not know is simply absent",
          "not-a-sample" not in res["counterparts"])
    check("the pane can name the file it consulted",
          res["metadata_file"].endswith("H5_metadata.xlsx") and res["metadata_rows"] == 2)

    hits = M.ref_find_sample("H5N1", q=on_disk, limit=50)
    check("the search tries the metadata spelling too",
          hits["also_searched"] == [in_table])
    check("...and finds the row", any(h["file"] == "H5_metadata.xlsx" for h in hits["hits"]))
    check("the stored name is found in column A with its .vcf suffix",
          any(h["text"] == f"{on_disk}_zc.vcf" for h in hits["hits"]))
    searched_by = M.ref_find_sample("H5N1", q=in_table, limit=50)
    check("searching the metadata name finds the stored one, marked as such",
          any(h["via_metadata"] and h["text"].startswith(on_disk)
              for h in searched_by["hits"]))
    none = M.ref_find_sample("H5N1", q="ZZZ-nothing", limit=50)
    check("a miss is a miss, with the files it looked in named",
          not none["hits"]
          and {x["file"] for x in none["searched"]}
              == {"H5_metadata.xlsx", "H5_define_filter.xlsx"})
    check("a workbook that would not open is reported, not counted as searched",
          [x["file"] for x in none["unsearchable"]] == ["H5_define_filter.xlsx"])

    print("[every other surface that matches a sample name]")
    # One project: a Step 1 sample renamed by the metadata, a panel sample
    # renamed by the GUI's own tree relabelling, and a remove list that holds
    # a metadata name by mistake.
    proj2 = tmp / "projects/p2"
    (proj2 / "step1/EPI-ISL-19152935/alignment_H5").mkdir(parents=True)
    (proj2 / "step1/ERR015582-trim200/alignment_H5").mkdir(parents=True)
    (proj2 / "step2/vcf_database").mkdir(parents=True)
    (proj2 / "step2/vcf_database/EPI-ISL-19152935_zc.vcf").write_text("")
    (proj2 / "step2/vcf_database/ERR015582-trim200_zc.vcf").write_text("")
    (proj2 / "project.json").write_text('{"reference": "H5N1"}', encoding="utf-8")
    # The csv names the ACCESSION (`ERR015582`); the relabel script rewrites it
    # wherever it stands as a whole word, so the tip for the sample stored as
    # `ERR015582-trim200` reads `Orygis_ERR015582-trim200` — observed on a real
    # labelled tree. Neither spelling is in any file.
    csv_path = tmp / "VCF_refs.csv"
    csv_path.write_text("M. orygis,ERR015582\n", encoding="utf-8")
    M._project_dir_for = lambda _cfg, _name: proj2
    M._find_vcf_refs_csv = lambda _cfg: csv_path
    M._PROJECT_ALIAS_CACHE.clear()

    idx2 = M._project_aliases(proj2, M.load_config())
    expected_tips = {f"{M._load_vcf_label_map(M.load_config(), style)['ERR015582']}_ERR015582-trim200"
                     for style in ("short", "rich")}
    check("the project index folds in the GUI's tree labels, both styles, "
          "as the script actually writes them",
          "Orygis_ERR015582-trim200" in expected_tips
          and all(idx2.original_of(t) == "ERR015582-trim200" for t in expected_tips))
    # Undone per lookup from the csv alone — no listing of the project's
    # samples, so nothing to rebuild when a Step 1 batch adds a folder. The
    # price is that a label resolves whether or not the project holds the
    # accession; every consumer checks the answer against what exists.
    check("a label resolves to its accession without consulting the project",
          idx2.original_of("Orygis_ERR015582") == "ERR015582")
    check("...and the accession standing mid-name is undone too",
          idx2.original_of("run7-Orygis_ERR015582") == "run7-ERR015582")
    check("a name that carries no label is not touched",
          idx2.original_of("ERR015582-trim200") is None)
    check("the forward direction gives every style's label",
          set(idx2.counterparts("ERR015582-trim200")) >= expected_tips)
    check("...and still the metadata", idx2.original_of(in_table) == on_disk)
    check("a clade selected on the relabelled tree resolves to the table's row",
          xlsx_html._selection_key("Orygis_ERR015582-trim200", {"ERR015582-trim200"},
                                   aliases=idx2) == "ERR015582-trim200")

    print("[_resolve_sample_dir: the resolver behind Step 1 files, IGV and the VCF editor]")
    step1 = proj2 / "step1"
    check("a stored name resolves as before",
          M._resolve_sample_dir(step1, "EPI-ISL-19152935") == step1 / "EPI-ISL-19152935")
    check("the metadata name now resolves to the same folder",
          M._resolve_sample_dir(step1, in_table) == step1 / "EPI-ISL-19152935")
    check("the tree label resolves too",
          M._resolve_sample_dir(step1, "Orygis_ERR015582-trim200") == step1 / "ERR015582-trim200")
    check("an unknown name is still None, not an error",
          M._resolve_sample_dir(step1, "nobody") is None)

    print("[/name-aliases: what the browser-side filters get]")
    na = M.project_name_aliases("p2")
    check("one entry per project sample that has another name, listing all of them",
          na["aliases"].get("EPI-ISL-19152935") == [in_table]
          and set(na["aliases"].get("ERR015582-trim200", [])) == expected_tips)
    check("...never the sample's own name",
          all(k not in v for k, v in na["aliases"].items()))

    print("[resolve-names: the tree search finds a relabelled tip]")
    rn = M.resolve_names("p2", M.NameResolveRequest(names=["ERR015582-trim200"]))
    check("the label is among the counterparts",
          "Orygis_ERR015582-trim200" in rn["counterparts"].get("ERR015582-trim200", []))

    print("[remove_from_analysis: vsnp3 matches FILE names, so a metadata name removes nothing]")
    import openpyxl as _ox
    rm = ref_dir / "H5_remove_from_analysis.xlsx"
    wb = _ox.Workbook(); ws = wb.active
    ws.cell(row=1, column=1, value="SRR1791698")          # a file stem: effective
    ws.cell(row=2, column=1, value=in_table)              # the SNP-table spelling: not
    wb.save(rm)
    state = M._remove_list_entries(ref_dir)
    check("the stem entry is effective",
          [e for e in state["entries"] if e["name"] == "SRR1791698"][0]["effective"] is True)
    check("the metadata-name entry is flagged, with the name vsnp3 needs",
          state["ineffective"] == [{"name": in_table, "stored": on_disk}])
    M.reference_roots = lambda _p: [ref_dir.parent]
    bl = M.step2_blocklist_get("p2")
    check("/step2/blocklist lists the file as written and names what will not work",
          "SRR1791698" in bl["samples"] and bl["ineffective"] == [{"name": in_table, "stored": on_disk}])

    M._ref_dir_or_404 = lambda _name: ref_dir
    M._current_os_user = lambda: "tester"
    M._T39_SHARED_AUDIT_PATH = tmp / "audit/reference-changes.jsonl"
    added = M.ref_remove_add_sample("H5N1", M.RemoveSampleAddRequest(
        samples=[in_table, "SRR1791698_updated_name", "ERR999"], rationale="test"))
    check("names pasted from a SNP table are written as the file names vsnp3 matches",
          added["translated"] == {in_table: on_disk, "SRR1791698_updated_name": "SRR1791698"})
    check("...the stem already listed is skipped, the new ones added",
          set(added["added"]) == {on_disk, "ERR999"} and added["skipped"] == ["SRR1791698"])
    fixed = M.ref_remove_normalize("H5N1", M.RemoveNormalizeRequest(rationale="test"))
    check("normalize rewrites the metadata-name entry",
          fixed["changed"] == [[in_table, on_disk]])
    names_after = M._read_remove_xlsx_names(rm)
    check("...leaving each sample listed once, as a file name",
          sorted(n for n in names_after if n) == sorted(["SRR1791698", on_disk, "ERR999"])
          and not M._remove_list_entries(ref_dir)["ineffective"])
    check("the previous file was archived first",
          Path(fixed["archived_old"]).is_file())

    print("[the Step 1 Results export keeps the rows the filtered table shows]")
    rows = [{"_sample": on_disk, "_run_date": "2026-09-01"},
            {"_sample": "ERR015582-trim200", "_run_date": "2026-09-01"},
            {"_sample": "SRR999", "_run_date": "2026-09-01"}]
    kept = M._qc_apply_filters(rows, None, None, "GISAID-19152935", aliases=idx2)
    check("a metadata name typed into the filter keeps the sample's row",
          [r["_sample"] for r in kept] == [on_disk])
    kept = M._qc_apply_filters(rows, None, None, "orygis", aliases=idx2)
    check("...and a tree label finds the panel sample", [r["_sample"] for r in kept] == ["ERR015582-trim200"])
    check("the stored name still matches on its own",
          [r["_sample"] for r in M._qc_apply_filters(rows, None, None, "srr999", aliases=idx2)] == ["SRR999"])
    check("no index, no change in behaviour",
          M._qc_apply_filters(rows, None, None, "GISAID") == [])

    print("PASS" if not FAILED else f"{FAILED} FAILURE(S)")
    return 1 if FAILED else 0


if __name__ == "__main__":
    raise SystemExit(main())
