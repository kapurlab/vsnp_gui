"""A Reference Editor edit must not delete the columns it does not understand.

Reported case: a lab opened a reference's ``*_metadata.xlsx`` and found two
columns. They keep submitter names, review notes and repeat-run dates in the
columns past B; adding one sample from the GUI had rewritten the workbook out
of the only two columns the editor reads, and everything else was gone. There
was no archived copy either — the metadata endpoint was the one reference-file
edit that skipped the ``.history/`` backup every other one takes.

What is pinned here: each mutating Reference Editor endpoint edits the sheet in
place, so cells outside its own two (or one) columns survive, as do other
sheets and rows below the last name; an append lands past the last row holding
ANYTHING, not past the last row holding a name; every edit archives first and
lands in the audit log; and the guarded save is real — a mutator that shrinks
the file is refused and the original put back.

The shared `taxa.yaml` is the same doctrine in a text file: adding a taxon
appends a line rather than rewriting the list, so the comments people keep
beside an entry are still there afterwards.

Run directly:  <conda>/bin/python backend/app/test_ref_edit_preserves_columns.py
"""
from __future__ import annotations

import hashlib
import json
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

import openpyxl  # noqa: E402
from fastapi import HTTPException  # noqa: E402

import main as M  # noqa: E402

FAILED = 0


def check(label, cond):
    global FAILED
    if cond:
        print(f"  OK  {label}")
    else:
        FAILED += 1
        print(f"  FAIL {label}")


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def cells(path: Path, sheet=0):
    ws = openpyxl.load_workbook(path).worksheets[sheet]
    return {(c.row, c.column): c.value
            for row in ws.iter_rows() for c in row if c.value is not None}


def main() -> int:
    tmp = Path(tempfile.mkdtemp(prefix="ref_edit_cols_"))
    ref_dir = tmp / "references" / "H5N1"
    ref_dir.mkdir(parents=True)

    M.load_config = lambda: {"vsnp3_path": str(tmp / "vsnp3")}
    M.list_references = lambda _p: [{"name": "H5N1", "path": str(ref_dir)}]
    M._ref_dir_or_404 = lambda _name: ref_dir
    M._current_os_user = lambda: "tester"
    M._T39_SHARED_AUDIT_PATH = tmp / "audit" / "reference-changes.jsonl"

    # ---------------------------------------------------------------- metadata
    print("[metadata: the lab's own columns are not ours to rewrite]")
    meta = ref_dir / "H5N1_metadata.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "metadata"
    for r, row in enumerate([
        ["sample", "display", "submitter", "notes"],
        ["EPI-ISL-1", "GISAID-1", "Vivek", "repeat run 2026-03"],
        ["EPI-ISL-2", "GISAID-2", "Tod", "low coverage, see log"],
        [None, None, None, "rows above reviewed 2026-09-01"],
    ], start=1):
        for c, v in enumerate(row, start=1):
            if v is not None:
                ws.cell(row=r, column=c, value=v)
    wb.create_sheet("lab notes")["A1"] = "primer lot 44"
    wb.save(meta)
    before_cells = cells(meta)
    before_sha = sha(meta)

    out = M.ref_add_metadata_rows("H5N1", M.MetadataAddRequest(rows=[
        {"original": "EPI-ISL-2", "display_name": "GISAID-2-corrected"},
        {"original": "EPI-ISL-9", "display_name": "GISAID-9"},
    ]))
    after = cells(meta)
    book = openpyxl.load_workbook(meta)

    check("the note columns are still there, cell for cell",
          all(after.get(k) == v for k, v in before_cells.items()
              if k[1] >= 3))
    check("the header row survived", after[(1, 1)] == "sample" and after[(1, 4)] == "notes")
    check("the display name of an existing sample was corrected in place",
          after[(3, 2)] == "GISAID-2-corrected")
    check("...without touching that row's submitter or note",
          after[(3, 3)] == "Tod" and after[(3, 4)] == "low coverage, see log")
    check("the trailing note-only row was not written over",
          after[(4, 4)] == "rows above reviewed 2026-09-01" and (4, 1) not in after)
    check("the new sample was appended BELOW it",
          after[(5, 1)] == "EPI-ISL-9" and after[(5, 2)] == "GISAID-9")
    check("the second sheet is still in the book",
          "lab notes" in book.sheetnames and book["lab notes"]["A1"].value == "primer lot 44")
    check("the sheet did not get narrower",
          max(c for _, c in after) == 4)
    check("the counts reported are the counts made",
          out["added"] == 1 and out["updated"] == 1 and out["rows_total"] == 4)

    print("[metadata: the edit is recoverable, like every other reference edit]")
    check("the previous file was archived first",
          out["archived_old"] and Path(out["archived_old"]).is_file()
          and sha(Path(out["archived_old"])) == before_sha)
    entries = [json.loads(l) for l in Path(out["audit_log"]).read_text().splitlines()]
    check("the edit is in the audit log, naming what changed",
          entries[-1]["action"] == "metadata_add_rows"
          and entries[-1]["added"] == ["EPI-ISL-9"]
          and entries[-1]["old_sha256"] == before_sha)

    print("[metadata: a re-read of the sheet says what it holds]")
    got = M.ref_get_metadata("H5N1")
    check("every name is listed once", len(got["rows"]) == 4)
    ws2 = openpyxl.load_workbook(meta).worksheets[0]
    ws2.cell(row=6, column=1, value="EPI-ISL-blank")
    ws2.cell(row=6, column=4, value="display name pending")
    ws2.parent.save(meta)
    blank = [r for r in M.ref_get_metadata("H5N1")["rows"]
             if r["original"] == "EPI-ISL-blank"][0]
    check("a row with no display name reads back blank, not the text 'nan'",
          blank["display_name"] == "")

    # ------------------------------------------------- remove_from_analysis
    print("[remove_from_analysis: the same file is shared, and so are its margins]")
    rm = ref_dir / "H5N1_remove_from_analysis.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.cell(row=1, column=1, value="SRR1")
    ws.cell(row=1, column=2, value="dropped: contaminated")
    ws.cell(row=2, column=2, value="everything below is pending review")
    wb.save(rm)
    rm_before = sha(rm)

    added = M.ref_remove_add_sample("H5N1", M.RemoveSampleAddRequest(
        samples=["SRR2"], rationale="test"))
    rm_after = cells(rm)
    check("the reason column beside the existing name survived",
          rm_after[(1, 2)] == "dropped: contaminated")
    check("the new name did NOT land on the row holding a standing note",
          rm_after[(2, 2)] == "everything below is pending review"
          and (2, 1) not in rm_after)
    check("...it went below it", rm_after[(3, 1)] == "SRR2" and added["added"] == ["SRR2"])
    check("that edit was archived too",
          sha(Path(added["archived_old"])) == rm_before)

    # ----------------------------------------------------------- the guard
    print("[the guarded save: a mutator that shrinks the file is refused]")
    keep = M._METADATA_ADD_CODE
    # Exactly the old bug: rebuild the sheet from the two columns we read.
    M._METADATA_ADD_CODE = M._XLSX_GUARDED_SAVE + r"""
import openpyxl, json, sys
target = sys.argv[1]
wb = openpyxl.load_workbook(target)
before = shape_of(wb)
pairs = [[r[0].value, r[1].value] for r in wb.worksheets[0].iter_rows(max_col=2)]
out = openpyxl.Workbook()
for i, (a, b) in enumerate(pairs, start=1):
    out.worksheets[0].cell(row=i, column=1, value=a)
    out.worksheets[0].cell(row=i, column=2, value=b)
guarded_save(out, target, before)
print(json.dumps({'added': [], 'updated': [], 'unchanged': [], 'rows_total': 0}))
"""
    doomed_sha = sha(meta)
    try:
        M.ref_add_metadata_rows("H5N1", M.MetadataAddRequest(
            rows=[{"original": "EPI-ISL-X", "display_name": "X"}]))
        refused = False
    except HTTPException as exc:
        refused = "shrunk" in str(exc.detail)
    finally:
        M._METADATA_ADD_CODE = keep
    check("the endpoint reports the refusal instead of the loss", refused)
    check("...and the workbook on disk is byte-for-byte the one we started with",
          sha(meta) == doomed_sha)

    print("[a workbook openpyxl cannot round-trip is refused, not flattened]")
    import zipfile
    charted = ref_dir / "H5N1_charted_metadata.xlsx"
    wb = openpyxl.Workbook()
    wb.active["A1"] = "S1"
    wb.active["B1"] = "display-1"
    wb.save(charted)
    # Splice in a chart part, the way a workbook saved from Excel carries one.
    tmp_zip = charted.with_suffix(".tmp.xlsx")
    with zipfile.ZipFile(charted) as zin, zipfile.ZipFile(tmp_zip, "w") as zout:
        for item in zin.infolist():
            zout.writestr(item, zin.read(item.filename))
        zout.writestr("xl/charts/chart1.xml", "<chart/>")
    tmp_zip.replace(charted)
    parked = tmp / "parked_metadata.xlsx"
    (ref_dir / "H5N1_metadata.xlsx").rename(parked)   # make the charted one the only match
    charted_sha = sha(charted)
    try:
        M.ref_add_metadata_rows("H5N1", M.MetadataAddRequest(
            rows=[{"original": "S2", "display_name": "display-2"}]))
        refused_chart = False
    except HTTPException as exc:
        refused_chart = "charts" in str(exc.detail)
    check("the edit is refused, naming what could not be preserved", refused_chart)
    check("...and the workbook is untouched", sha(charted) == charted_sha)
    charted.unlink()
    parked.rename(ref_dir / "H5N1_metadata.xlsx")

    print("[taxa.yaml: adding a name does not rewrite the file over its comments]")
    taxa = tmp / "taxa.yaml"
    taxa.write_text(
        "# Kraken ID Parse — taxon search names\n"
        "#\n"
        "# add by hand or via the GUI control.\n"
        "\n"
        "- Mycobacterium tuberculosis     # MTBC0 reference set\n"
        "# added for the 2026 outbreak — remove after the season\n"
        "- Brucella abortus\n"
    )
    M._KRAKEN_TAXA_YAML = taxa
    M._append_kraken_taxon("Salmonella enterica")
    text = taxa.read_text()
    check("a standalone comment survived",
          "# added for the 2026 outbreak — remove after the season" in text)
    check("an inline comment survived", "# MTBC0 reference set" in text)
    check("the new name is on the end", text.rstrip().endswith("- Salmonella enterica"))
    check("and every name reads back",
          M._read_kraken_taxa() == ["Mycobacterium tuberculosis", "Brucella abortus",
                                    "Salmonella enterica"])

    print("PASS" if not FAILED else f"{FAILED} FAILURE(S)")
    return 1 if FAILED else 0


if __name__ == "__main__":
    raise SystemExit(main())
