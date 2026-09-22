"""A reference entry that is a symlink stays one, and keeps its permissions.

A curated reference tree keeps one spreadsheet and links it into every
reference that uses it — on the USDA Ames HPC one `FileMakerTB_metadata.xlsx`
under `vsnp_dependencies/` serves a dozen references. Two things went wrong
when the GUI wrote such an entry:

  * Replace installed the upload by RENAME, which put a new file at the link's
    own path. The reference stopped sharing without saying so: the master kept
    the old content and every sibling reference kept reading it.
  * A rename also puts a NEW INODE at the path, so everything about the file
    that was not its content went with the old one — owner, mode, ACL. A
    workbook the whole group could read came back `-rw-------` owned by the
    uploader, and the other users of that reference could no longer read it.

Both are fixed the same way: find the file the entry actually points at, check
it can be written, and write THROUGH it instead of over it.

Run directly:  <conda>/bin/python backend/app/test_ref_edit_follows_symlinks.py
"""
from __future__ import annotations

import asyncio
import io
import json
import os
import stat
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

import openpyxl  # noqa: E402
from fastapi import HTTPException, UploadFile  # noqa: E402

import main as M  # noqa: E402

FAILED = 0


def check(label, cond):
    global FAILED
    if cond:
        print(f"  OK  {label}")
    else:
        FAILED += 1
        print(f"  FAIL {label}")


def names_in(path: Path):
    ws = openpyxl.load_workbook(path).worksheets[0]
    return [ws.cell(row=r, column=1).value for r in range(1, ws.max_row + 1)]


def notes_in(path: Path):
    ws = openpyxl.load_workbook(path).worksheets[0]
    return [ws.cell(row=r, column=4).value for r in range(1, ws.max_row + 1)]


def make_master(path: Path):
    wb = openpyxl.Workbook()
    ws = wb.active
    for r, row in enumerate([
        ["sample", "display", "submitter", "notes"],
        ["EPI-1", "GISAID-1", "kalehman", "curated 2026-08"],
    ], start=1):
        for c, v in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=v)
    wb.save(path)


def upload(ref_name: str, filename: str, data: bytes, rationale="test"):
    f = UploadFile(file=io.BytesIO(data), filename=filename)
    return asyncio.run(M.ref_upload_file(ref_name, file=f, rationale=rationale))


def main() -> int:
    tmp = Path(tempfile.mkdtemp(prefix="ref_symlink_"))
    deps = tmp / "vsnp_dependencies"
    ref_a = deps / "Mycobacterium_AF2122"
    ref_b = deps / "Mycobacterium_H37"
    for d in (ref_a, ref_b):
        d.mkdir(parents=True)

    shared = deps / "FileMakerTB_metadata.xlsx"
    make_master(shared)
    # macOS puts temp dirs under /var, itself a link to /private/var, so the
    # canonical path the resolver returns is not the one we built.
    real_shared = os.path.realpath(shared)
    os.chmod(shared, 0o664)
    for d in (ref_a, ref_b):
        (d / "FileMakerTB_metadata.xlsx").symlink_to(shared)

    REFS = [{"name": "AF2122", "path": str(ref_a)}, {"name": "H37", "path": str(ref_b)}]
    M.load_config = lambda: {"vsnp3_path": str(tmp / "vsnp3")}
    M.list_references = lambda _p: REFS
    M._ref_dir_or_404 = lambda name: ref_a if name == "AF2122" else ref_b
    M._current_os_user = lambda: "tstuber"
    M._T39_SHARED_AUDIT_PATH = tmp / "audit" / "reference-changes.jsonl"

    link_a = ref_a / "FileMakerTB_metadata.xlsx"
    ino = shared.stat().st_ino
    mode = stat.S_IMODE(shared.stat().st_mode)

    print("[add a row through a link: the master is what changes]")
    out = M.ref_add_metadata_rows("AF2122", M.MetadataAddRequest(
        rows=[{"original": "EPI-2", "display_name": "GISAID-2"}]))
    check("the entry is still a symlink", link_a.is_symlink())
    check("...still pointing at the master", os.path.realpath(link_a) == real_shared)
    check("the master is the same file, not a replacement",
          shared.stat().st_ino == ino)
    check("its mode is untouched", stat.S_IMODE(shared.stat().st_mode) == mode)
    check("the new row is in the master", "EPI-2" in names_in(shared))
    check("the curator's note column survived", "curated 2026-08" in notes_in(shared))
    check("the other reference sees the edit through its own link",
          "EPI-2" in names_in(ref_b / "FileMakerTB_metadata.xlsx"))

    print("[the record says where the edit landed and who else reads it]")
    check("the response names the master", out.get("master") == real_shared)
    check("...and the references that share it", out.get("also_used_by") == ["H37"])
    check("the archive is filed beside the master, not under one reference",
          Path(out["archived_old"]).parent == Path(real_shared).parent / ".history")
    check("the archived copy is the master's previous BYTES, not a link",
          not Path(out["archived_old"]).is_symlink()
          and "EPI-2" not in names_in(Path(out["archived_old"])))
    entry = json.loads(Path(out["audit_log"]).read_text().splitlines()[-1])
    check("the audit line carries the same facts",
          entry["master"] == real_shared and entry["also_used_by"] == ["H37"])

    print("[Replace through a link: same file, same permissions]")
    replacement = tmp / "replacement.xlsx"
    make_master(replacement)
    wb = openpyxl.load_workbook(replacement)
    wb.worksheets[0].cell(row=3, column=1, value="EPI-9")
    wb.worksheets[0].cell(row=3, column=4, value="added offline")
    wb.save(replacement)

    res = upload("AF2122", "FileMakerTB_metadata.xlsx", replacement.read_bytes())
    check("the entry is still a symlink", link_a.is_symlink())
    check("the master is still the same inode", shared.stat().st_ino == ino)
    check("...with its mode intact", stat.S_IMODE(shared.stat().st_mode) == mode)
    check("the uploaded content is in the master", "EPI-9" in names_in(shared))
    check("the sibling reference sees it too",
          "EPI-9" in names_in(ref_b / "FileMakerTB_metadata.xlsx"))
    check("the response names the master and its sharers",
          res.get("master") == real_shared and res.get("also_used_by") == ["H37"])
    check("the previous version was archived first",
          "EPI-2" in names_in(Path(res["archived_old"])))

    print("[Replace on an ordinary file keeps ITS inode and mode too]")
    plain = ref_b / "H37_remove_from_analysis.xlsx"
    wb = openpyxl.Workbook()
    wb.active.cell(row=1, column=1, value="SRR1")
    wb.save(plain)
    os.chmod(plain, 0o664)
    p_ino, p_mode = plain.stat().st_ino, stat.S_IMODE(plain.stat().st_mode)
    swap = tmp / "swap.xlsx"
    wb = openpyxl.Workbook()
    wb.active.cell(row=1, column=1, value="SRR2")
    wb.save(swap)
    upload("H37", "H37_remove_from_analysis.xlsx", swap.read_bytes())
    check("the file was written through, not replaced", plain.stat().st_ino == p_ino)
    check("...so its mode (and any ACL on it) survived",
          stat.S_IMODE(plain.stat().st_mode) == p_mode)
    check("and the new content is there", names_in(plain) == ["SRR2"])

    print("[a master this account cannot write is refused, not un-shared]")
    os.chmod(shared, 0o444)
    before = shared.stat().st_ino
    try:
        M.ref_add_metadata_rows("AF2122", M.MetadataAddRequest(
            rows=[{"original": "EPI-X", "display_name": "X"}]))
        refused_add = False
    except HTTPException as exc:
        refused_add = exc.status_code == 403 and real_shared in str(exc.detail)
    try:
        upload("AF2122", "FileMakerTB_metadata.xlsx", replacement.read_bytes())
        refused_replace = False
    except HTTPException as exc:
        refused_replace = exc.status_code == 403 and real_shared in str(exc.detail)
    check("the add is refused, naming the master", refused_add)
    check("the replace is refused, naming the master", refused_replace)
    check("the link is still a link, pointing where it did",
          link_a.is_symlink() and os.path.realpath(link_a) == real_shared)
    check("nothing was written", shared.stat().st_ino == before
          and "EPI-X" not in names_in(shared))
    os.chmod(shared, 0o664)

    print("PASS" if not FAILED else f"{FAILED} FAILURE(S)")
    return 1 if FAILED else 0


if __name__ == "__main__":
    raise SystemExit(main())
