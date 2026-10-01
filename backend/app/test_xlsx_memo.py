#!/usr/bin/env python3
"""The streaming renderers' memos answer exactly what the per-cell helpers do.

_CfMemo and _CellStyleMemo (xlsx_html.py) exist so a cold render of a wide
cascade table stops re-evaluating conditional formatting and style for every
one of its hundreds of thousands of cells. They are only allowed to be faster:
for every cell of a sheet built like vSNP3's — a reference row the top rule
compares against (`equal B$2`), one containsText rule per base, a numeric rule
on the MQ rows, a rule with a relative reference that can never resolve, and
values of mixed types that compare equal in Python (1, 1.0, True) — the memo
must return the very dxf _cf_dxf_for returns and the very strings the style
helpers return. And it must actually memoise: far fewer evaluations than cells.

Run: cd backend/app && ../../env/bin/python test_xlsx_memo.py
"""
from __future__ import annotations

import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

import openpyxl
from openpyxl.formatting.rule import CellIsRule, Rule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.utils import get_column_letter

import xlsx_html

FAILURES = []


def check(cond, label):
    print(("ok   " if cond else "FAIL ") + label)
    if not cond:
        FAILURES.append(label)


def make_table(path: Path, rows: int, cols: int) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.cell(row=1, column=1, value="Sample")
    for c in range(2, cols + 1):
        ws.cell(row=1, column=c, value=f"NC_000962:{1000 + c * 7}")
        ws.cell(row=1, column=c).alignment = Alignment(text_rotation=90)
    ws.cell(row=2, column=1, value="root")
    bases = ["A", "C", "G", "T"]
    for c in range(2, cols + 1):
        ws.cell(row=2, column=c, value=bases[c % 4])
    mixed = ["A", "C", "G", "T", "-", "N", "a", 1, 1.0, True, "1", 0, 0.0, False, "True", 55, 54.5, "55", None]
    for r in range(3, rows - 1):
        ws.cell(row=r, column=1, value=f"SAMPLE{r - 2}_zc.vcf")
        for c in range(2, cols + 1):
            v = mixed[(r * 3 + c) % len(mixed)]
            if v is not None:
                ws.cell(row=r, column=c, value=v)
            if (r + c) % 5 == 0:
                ws.cell(row=r, column=c).font = Font(bold=True, color="FF336699")
    ws.cell(row=rows - 1, column=1, value="MQ")
    ws.cell(row=rows, column=1, value="MQ2")
    for c in range(2, cols + 1):
        ws.cell(row=rows - 1, column=c, value=50 + (c % 12))
        ws.cell(row=rows, column=c, value=float(53 + (c % 5)) + 0.5)
    last = get_column_letter(cols)
    body = f"B3:{last}{rows - 2}"
    # vSNP3's rules: the reference match first, then one colour per base.
    ws.conditional_formatting.add(body, CellIsRule(
        operator="equal", formula=["B$2"],
        fill=PatternFill(start_color="FFF2F2F2", end_color="FFF2F2F2", fill_type="solid")))
    for base, colour in (("A", "FF00FF00"), ("G", "FFFF0000"), ("C", "FF0000FF"), ("T", "FFFFFF00"), ("N", "FF999999")):
        ws.conditional_formatting.add(body, Rule(
            type="containsText", operator="containsText", text=base,
            formula=[f'NOT(ISERROR(SEARCH("{base}",B3)))'],
            dxf=DifferentialStyle(fill=PatternFill(bgColor=colour, fill_type="solid"))))
    # A relative-row reference the streaming pass can never resolve.
    ws.conditional_formatting.add(f"B3:D{rows - 2}", CellIsRule(
        operator="notEqual", formula=["B2"],
        font=Font(italic=True, color="FF112233")))
    # The MQ rows: a numeric rule, the kind that compares as floats.
    ws.conditional_formatting.add(f"B{rows - 1}:{last}{rows}", CellIsRule(
        operator="lessThan", formula=["55"],
        fill=PatternFill(start_color="FFFFC000", end_color="FFFFC000", fill_type="solid")))
    wb.save(path)
    wb.close()


def main() -> int:
    tmp = Path(tempfile.mkdtemp(prefix="xlsx_memo_"))
    try:
        rows, cols = 40, 60
        book = tmp / "table.xlsx"
        make_table(book, rows, cols)
        wb = openpyxl.load_workbook(book, data_only=True, read_only=True)
        try:
            ws = wb.active
            dxfs = list(wb._differential_styles.styles)
            cf_ranges = xlsx_html._cf_from_sheet_xml(book)
            check(len(cf_ranges) >= 3, f"the test sheet carries several CF ranges ({len(cf_ranges)})")
            cf_sheet = xlsx_html._CapturedRowsSheet()
            capture_rows = xlsx_html._cf_absolute_ref_rows(cf_ranges)
            check(2 in capture_rows, "the reference row is the one captured for `equal B$2`")
            cf_memo = xlsx_html._CfMemo(cf_ranges, dxfs, cf_sheet, cols)
            style_memo = xlsx_html._CellStyleMemo()
            cells = 0
            dxf_mismatch = frag_mismatch = style_mismatch = 0
            coloured = 0
            for row_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=rows, min_col=1, max_col=cols), start=1):
                if row_idx in capture_rows:
                    cf_sheet.capture(row_idx, row)
                for col_idx, cell in enumerate(row, start=1):
                    if not hasattr(cell, "column"):
                        continue
                    cells += 1
                    want = xlsx_html._cf_dxf_for(cell, row_idx, col_idx, cf_ranges, dxfs, cf_sheet)
                    got = cf_memo.dxf_for(cell, row_idx, col_idx)
                    if got is not want:
                        dxf_mismatch += 1
                    if want is not None:
                        coloured += 1
                    want_frags = xlsx_html._cf_fragments_for(cell, row_idx, col_idx, cf_ranges, dxfs, cf_sheet)
                    if list(cf_memo.fragments_for(cell, row_idx, col_idx)) != want_frags:
                        frag_mismatch += 1
                    if (style_memo.inline_style(cell) != xlsx_html._cell_inline_style(cell)
                            or style_memo.rotation_class(cell) != xlsx_html._cell_rotation_class(cell)
                            or style_memo.value_html(cell) != xlsx_html._format_cell_value(cell)):
                        style_mismatch += 1
            check(cells > 2000, f"enough cells to mean something ({cells})")
            check(coloured > 500, f"the rules colour a good share of them ({coloured})")
            check(dxf_mismatch == 0, f"the CF memo returns the evaluator's dxf for every cell ({dxf_mismatch} differed)")
            check(frag_mismatch == 0, f"and the same CSS fragments ({frag_mismatch} differed)")
            check(style_mismatch == 0, f"the style memo returns the helpers' strings for every cell ({style_mismatch} differed)")
            evaluations = len(cf_memo._dxf_memo)
            check(evaluations < cells / 4, f"CF was evaluated {evaluations} times for {cells} cells")
            check(len(style_memo._inline) < 20, f"inline style computed {len(style_memo._inline)} times")
            # Type-tagged keys: 1, 1.0 and True are distinct entries, not one.
            types_seen = {k[1] for k in style_memo._values if k[2] == 1}
            check(len(types_seen) >= 2, f"1, 1.0 and True keep separate value entries ({sorted(t.__name__ for t in types_seen)})")
        finally:
            wb.close()

        # And through the renderers themselves: the page a table renders to is
        # the same whether or not the memos sit in the loop. The reference is
        # the streaming renderer with the memos bypassed (a stub that answers
        # from the helpers directly).
        samples = {f"SAMPLE{i}" for i in range(1, rows)}
        with_memo = xlsx_html.render_window(book, rows, cols, None, "p", samples, set(),
                                            xlsx_html.DEFAULT_MAX_CELLS, xlsx_html.DEFAULT_MAX_ROWS)

        class _NoCf:
            def __init__(self, cf_ranges, dxfs, cf_sheet, ncols):
                self.a = (cf_ranges, dxfs, cf_sheet)

            def dxf_for(self, cell, r, c):
                return xlsx_html._cf_dxf_for(cell, r, c, *self.a)

            def fragments_for(self, cell, r, c):
                return xlsx_html._cf_fragments_for(cell, r, c, *self.a)

        class _NoStyle:
            inline_style = staticmethod(xlsx_html._cell_inline_style)
            rotation_class = staticmethod(xlsx_html._cell_rotation_class)
            value_html = staticmethod(xlsx_html._format_cell_value)

        saved = (xlsx_html._CfMemo, xlsx_html._CellStyleMemo)
        xlsx_html._CfMemo, xlsx_html._CellStyleMemo = _NoCf, _NoStyle
        try:
            without = xlsx_html.render_window(book, rows, cols, None, "p", samples, set(),
                                              xlsx_html.DEFAULT_MAX_CELLS, xlsx_html.DEFAULT_MAX_ROWS)
            filtered_without = xlsx_html.render_filtered_window(
                book, rows, cols, None, "p", samples, set(), ["SAMPLE3", "SAMPLE7", "SAMPLE11"],
                xlsx_html.DEFAULT_MAX_CELLS, xlsx_html.DEFAULT_MAX_ROWS)
        finally:
            xlsx_html._CfMemo, xlsx_html._CellStyleMemo = saved
        filtered_with = xlsx_html.render_filtered_window(
            book, rows, cols, None, "p", samples, set(), ["SAMPLE3", "SAMPLE7", "SAMPLE11"],
            xlsx_html.DEFAULT_MAX_CELLS, xlsx_html.DEFAULT_MAX_ROWS)
        check(with_memo == without, "render_window: identical window with and without the memos")
        check(filtered_with == filtered_without, "render_filtered_window: identical window with and without the memos")
        check(sum(r.count("xlsx-variant") for r in with_memo["rows"]) > 100,
              "the rendered page has coloured variant cells (the CF took effect)")
    finally:
        import shutil
        shutil.rmtree(tmp, ignore_errors=True)
    if FAILURES:
        print(f"\n{len(FAILURES)} check(s) failed")
        return 1
    print("\nall checks passed")
    return 0


if __name__ == "__main__":
    sys.exit(main())
